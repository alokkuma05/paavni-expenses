"""Iteration 4 tests - Recurring Expenses + Vendor Statement + Budget Alerts."""
import os
import pytest
import requests
from io import BytesIO
from openpyxl import load_workbook

BASE_URL = os.environ["EXPO_PUBLIC_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def api_client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


# ---------------------------------------------------------------------------
# Recurring Expenses
# ---------------------------------------------------------------------------
class TestRecurringExpenses:
    """POST /api/expenses with recurring:true creates template + instance."""

    created_expense_ids = []
    created_template_ids = []

    def test_create_recurring_expense_creates_template(self, api_client):
        payload = {
            "date": "2026-10-07",
            "vendor": "TEST_RecurVendor",
            "main_category": "Office Expenses",
            "amount": 1500,
            "payment_status": "Unpaid",
            "recurring": True,
            "remarks": "TEST_recur_bill",
        }
        resp = api_client.post(f"{API}/expenses", json=payload)
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["recurring"] is True
        assert data.get("recurring_id"), "recurring_id must be set on created expense"
        TestRecurringExpenses.created_expense_ids.append(data["id"])
        TestRecurringExpenses.created_template_ids.append(data["recurring_id"])

        # Verify template exists and day_of_month=7 derived from the date
        resp = api_client.get(f"{API}/recurring")
        assert resp.status_code == 200
        templates = resp.json()
        tmpl = next((t for t in templates if t["id"] == data["recurring_id"]), None)
        assert tmpl is not None, "Template must be returned by GET /api/recurring"
        assert tmpl["active"] is True
        assert tmpl["day_of_month"] == 7
        assert tmpl["vendor"] == "TEST_RecurVendor"
        assert tmpl["amount"] == 1500.0

    def test_recurring_idempotency_no_duplicates(self, api_client):
        """GET /api/expenses twice must not duplicate the current month instance."""
        # First call runs ensure_recurring()
        r1 = api_client.get(f"{API}/expenses", params={"vendor": "Office Landlord", "month": "2026-10"})
        assert r1.status_code == 200
        count1 = len(r1.json())
        # Second call -- idempotent
        r2 = api_client.get(f"{API}/expenses", params={"vendor": "Office Landlord", "month": "2026-10"})
        assert r2.status_code == 200
        count2 = len(r2.json())
        assert count1 == count2, f"Office Landlord Oct count must stay stable, got {count1} -> {count2}"
        # Expect exactly 1 from pre-seed
        assert count2 == 1, f"Office Landlord should have exactly 1 Oct 2026 instance, got {count2}"

    def test_recurring_instance_linked_via_recurring_id(self, api_client):
        """The Office Landlord Oct 2026 instance should be linked to its template."""
        r = api_client.get(f"{API}/expenses", params={"vendor": "Office Landlord", "month": "2026-10"})
        assert r.status_code == 200
        bills = r.json()
        assert len(bills) >= 1
        bill = bills[0]
        assert bill["recurring"] is True
        assert bill.get("recurring_id"), "Auto-generated instance must have recurring_id"
        assert bill["payment_status"] == "Unpaid"
        assert bill["main_category"] == "Office Expenses"
        assert bill["amount"] == 25000

    def test_delete_recurring_soft_deletes(self, api_client):
        """DELETE /api/recurring/{id} sets active=false and removes from list."""
        if not TestRecurringExpenses.created_template_ids:
            pytest.skip("No test template created")
        tmpl_id = TestRecurringExpenses.created_template_ids[0]
        r = api_client.delete(f"{API}/recurring/{tmpl_id}")
        assert r.status_code == 200
        # Verify removed from active list
        r2 = api_client.get(f"{API}/recurring")
        assert r2.status_code == 200
        ids = [t["id"] for t in r2.json()]
        assert tmpl_id not in ids, "Deleted template must not appear in GET /api/recurring"

    def test_cleanup_recurring_expenses(self, api_client):
        """Cleanup test-created expenses (templates already soft-deleted above)."""
        for eid in TestRecurringExpenses.created_expense_ids:
            api_client.delete(f"{API}/expenses/{eid}")


# ---------------------------------------------------------------------------
# Vendor Statement
# ---------------------------------------------------------------------------
class TestVendorStatement:
    """GET /api/expenses?vendor and GET /api/export?vendor."""

    def test_filter_expenses_by_vendor(self, api_client):
        r = api_client.get(f"{API}/expenses", params={"vendor": "Office Landlord"})
        assert r.status_code == 200
        bills = r.json()
        assert len(bills) >= 1
        for b in bills:
            assert b["vendor"] == "Office Landlord", f"Expected only Office Landlord, got {b['vendor']}"

    def test_export_vendor_returns_xlsx(self, api_client):
        r = api_client.get(f"{API}/export", params={"vendor": "Office Landlord"})
        assert r.status_code == 200
        ct = r.headers.get("content-type", "")
        assert "spreadsheet" in ct or "xlsx" in ct, f"Expected xlsx content-type, got {ct}"
        cd = r.headers.get("content-disposition", "")
        assert "office-landlord" in cd.lower(), f"Vendor slug must be in filename: {cd}"
        # Verify it's a valid xlsx
        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        # Header row + at least 1 data row + total row
        assert ws.max_row >= 3
        # Verify every data row is for Office Landlord (col B is Vendor)
        for row in range(2, ws.max_row):
            v = ws.cell(row=row, column=2).value
            # The total row has empty vendor
            if v:
                assert v == "Office Landlord", f"Row {row} vendor should be Office Landlord, got {v}"


# ---------------------------------------------------------------------------
# Budget Alerts (dashboard pct)
# ---------------------------------------------------------------------------
class TestBudgetAlerts:
    """GET /api/dashboard?month includes pct for each budget category."""

    def test_dashboard_budget_block_has_pct(self, api_client):
        r = api_client.get(f"{API}/dashboard", params={"month": "2026-05"})
        assert r.status_code == 200
        data = r.json()
        budget = data.get("budget") or {}
        assert budget.get("has_budgets") is True, "May 2026 should have budgets"
        by_cat = budget.get("by_category") or []
        assert len(by_cat) >= 1
        for item in by_cat:
            assert "pct" in item
            assert "budget" in item
            assert "spent" in item
            assert "remaining" in item
            assert "category" in item

    def test_milk_procurement_cost_110_pct(self, api_client):
        """Seeded: Milk Procurement Cost budget 50000, actual 55000 => 110%."""
        r = api_client.get(f"{API}/dashboard", params={"month": "2026-05"})
        assert r.status_code == 200
        by_cat = r.json()["budget"]["by_category"]
        mpc = next((b for b in by_cat if b["category"] == "Milk Procurement Cost"), None)
        assert mpc is not None, "Milk Procurement Cost budget must exist for 2026-05"
        assert mpc["budget"] == 50000
        assert mpc["spent"] == 55000
        assert mpc["pct"] == 110.0, f"Expected pct 110.0, got {mpc['pct']}"
        # Verify this is flaggable at >=90%
        assert mpc["pct"] >= 90.0, "Category should be flaggable as >=90%"
