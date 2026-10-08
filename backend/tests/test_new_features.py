"""Backend tests for the 4 new features: payment ledger, vendor dues, Excel export, product_line."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("EXPO_PUBLIC_BACKEND_URL", "https://expense-tracker-7630.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def api_client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def created():
    ids = {"expenses": []}
    yield ids
    for eid in ids["expenses"]:
        try:
            requests.delete(f"{API}/expenses/{eid}")
        except Exception:
            pass


# ---- Payment ledger ----
class TestPaymentLedger:
    def test_create_unpaid_has_empty_ledger(self, api_client, created):
        r = api_client.post(f"{API}/expenses", json={
            "date": "2026-03-01", "vendor": "TEST_Ledger_A",
            "main_category": "Production Cost", "amount": 1000,
            "paid_amount": 0, "payment_status": "Unpaid",
            "product_line": "Paneer",
        })
        assert r.status_code == 200
        e = r.json()
        created["expenses"].append(e["id"])
        assert e["payments"] == []
        assert e["product_line"] == "Paneer"
        assert e["outstanding"] == 1000

    def test_create_with_initial_payment_seeds_ledger(self, api_client, created):
        r = api_client.post(f"{API}/expenses", json={
            "date": "2026-03-02", "vendor": "TEST_Ledger_B",
            "main_category": "Office Expenses", "amount": 1000,
            "paid_amount": 300, "payment_status": "Unpaid",
            "payment_mode": "UPI", "payment_date": "2026-03-02",
            "product_line": "Curd",
        })
        assert r.status_code == 200
        e = r.json()
        created["expenses"].append(e["id"])
        assert len(e["payments"]) == 1
        p = e["payments"][0]
        assert p["amount"] == 300
        assert p["mode"] == "UPI"
        assert p["note"] == "Initial payment"
        assert e["payment_status"] == "Partly Paid"
        assert e["outstanding"] == 700

    def test_record_payment_appends_and_updates_status(self, api_client, created):
        eid = created["expenses"][0]  # 1000/0 Unpaid
        r = api_client.post(f"{API}/expenses/{eid}/payment", json={
            "paid_amount": 400, "payment_mode": "Cash", "payment_date": "2026-03-05", "note": "First part"
        })
        assert r.status_code == 200
        e = r.json()
        assert e["paid_amount"] == 400
        assert e["outstanding"] == 600
        assert e["payment_status"] == "Partly Paid"
        assert len(e["payments"]) == 1
        assert e["payments"][0]["note"] == "First part"

        # Second part - complete it
        r = api_client.post(f"{API}/expenses/{eid}/payment", json={
            "paid_amount": 600, "payment_mode": "Bank Transfer", "payment_date": "2026-03-07"
        })
        assert r.status_code == 200
        e = r.json()
        assert e["paid_amount"] == 1000
        assert e["outstanding"] == 0
        assert e["payment_status"] == "Paid"
        assert len(e["payments"]) == 2
        assert e["payments"][1]["amount"] == 600
        assert e["payments"][1]["mode"] == "Bank Transfer"

        # GET to verify persistence
        g = api_client.get(f"{API}/expenses/{eid}")
        assert g.status_code == 200
        assert len(g.json()["payments"]) == 2

    def test_record_payment_rejects_zero(self, api_client, created):
        eid = created["expenses"][1]
        r = api_client.post(f"{API}/expenses/{eid}/payment", json={"paid_amount": 0})
        assert r.status_code == 400

    def test_record_payment_rejects_negative(self, api_client, created):
        eid = created["expenses"][1]
        r = api_client.post(f"{API}/expenses/{eid}/payment", json={"paid_amount": -50})
        assert r.status_code == 400

    def test_record_payment_404(self, api_client):
        r = api_client.post(f"{API}/expenses/does-not-exist/payment", json={"paid_amount": 100})
        assert r.status_code == 404


# ---- Vendor dues ----
class TestVendorDues:
    def test_dues_shape_and_sort(self, api_client, created):
        # created so far: Ledger_A (1000/1000, out 0), Ledger_B (1000/300, out 700)
        r = api_client.get(f"{API}/vendors/dues")
        assert r.status_code == 200
        d = r.json()
        assert "vendors" in d and "total_outstanding" in d
        # Sorted desc by outstanding
        outs = [v["outstanding"] for v in d["vendors"]]
        assert outs == sorted(outs, reverse=True)
        # Each vendor has required keys
        for v in d["vendors"]:
            for k in ["vendor", "total", "paid", "outstanding", "bills", "open_bills"]:
                assert k in v
        # TEST_Ledger_B should be present with outstanding=700
        b = next((v for v in d["vendors"] if v["vendor"] == "TEST_Ledger_B"), None)
        assert b is not None
        assert b["outstanding"] == 700
        assert b["open_bills"] == 1
        assert b["bills"] == 1
        # total_outstanding is sum of outstanding
        assert round(d["total_outstanding"], 2) == round(sum(v["outstanding"] for v in d["vendors"]), 2)


# ---- Excel export ----
class TestExport:
    def test_export_all(self, api_client):
        r = api_client.get(f"{API}/export")
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower()
        assert ".xlsx" in cd
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # header + at least 1 data row + TOTAL row
        assert ws.max_row >= 3
        headers = [c.value for c in ws[1]]
        assert "Date" in headers and "Amount" in headers and "Product Line" in headers
        # TOTAL row is last
        last = [c.value for c in ws[ws.max_row]]
        assert "TOTAL" in last

    def test_export_month(self, api_client):
        r = api_client.get(f"{API}/export", params={"month": "2026-03"})
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        assert ws.title == "2026-03"
        # should contain our TEST_Ledger_A/B rows
        all_rows = []
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 1, values_only=True):
            all_rows.append(row)
        vendors = [row[1] for row in all_rows]
        assert "TEST_Ledger_A" in vendors or "TEST_Ledger_B" in vendors
        # Product Line column (index 6) populated for TEST_Ledger_B = Curd
        for row in all_rows:
            if row[1] == "TEST_Ledger_B":
                assert row[6] == "Curd"


# ---- product_line persistence on update ----
class TestProductLine:
    def test_update_persists_product_line(self, api_client, created):
        eid = created["expenses"][1]  # TEST_Ledger_B
        r = api_client.put(f"{API}/expenses/{eid}", json={
            "date": "2026-03-02", "vendor": "TEST_Ledger_B",
            "main_category": "Office Expenses", "amount": 1000,
            "paid_amount": 300, "payment_status": "Partly Paid",
            "product_line": "Ghee",
        })
        assert r.status_code == 200
        assert r.json()["product_line"] == "Ghee"
        g = api_client.get(f"{API}/expenses/{eid}")
        assert g.json()["product_line"] == "Ghee"
