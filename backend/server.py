from fastapi import FastAPI, APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import calendar
import logging
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional
import uuid
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Master data (seeded from the Paavni Monthly Expenses sheet)
# ---------------------------------------------------------------------------
CATEGORY_HIERARCHY = {
    "Production Cost": {
        "Cattle Feed": ["Dry Fodder", "Silage", "Green Fodder", "Cattle Feed Supplement", "Other Supplements", "Fodder"],
        "Salary": ["Farm Manager Payroll", "Farm Supervisor Payroll", "Permanent Labor Payroll", "Casual Labour Payroll"],
        "Veterinary": ["Doctors Fees", "Medicines", "Insurance Fees", "Veterinary Exp"],
        "Machinery Repair": ["Milking Machine Repair & Maintenance", "Gen Set Repair", "Electricity Repair", "Other Farm Repair"],
        "Energy Cost": ["Electricity Bill", "Gen Set Diesel", "Gen Set Rent", "Bike Petrol", "Fuel"],
    },
    "Milk Procurement Cost": {
        "Milk Purchase Cost": ["Purchase of Milk", "Milk Purchase Cost"],
        "Energy Cost": ["CNG Fuel & Transportation", "Diesel Fuel & Transportation", "Fuel - Milk Procurement"],
        "Salary": ["Procurement Supervisor Payroll", "Procurement Labour Payroll"],
    },
    "Milk Processing and Packaging": {
        "Energy Cost": ["Electricity Bill - Processing", "Diesel - Processing", "LPG"],
        "Packaging": ["Pouch Purchase", "Bottle Purchase", "Bottle Sealing", "Bottle Cleaning", "Packaging Materials"],
        "Salary": ["Processing Manager Payroll", "Processing Supervisor Payroll", "Processing Labour Payroll", "Processing Casual Labour Payroll"],
        "Machinery Repair & Maintenance": ["BMC Repair", "BMC Cleaning", "Processing Unit Repair", "Technician Fees", "Depreciation"],
        "Rent of Premises": ["Processing Unit Rent", "Rent Farm"],
    },
    "Milk Transportation Cost": {
        "Vehicle Fuel": ["Fuel - Milk Transportation", "Fast Tag/Toll Charges", "Vehicle Repair", "Vehicle Rent", "Vehicle Rent/EMI"],
        "Salary": ["Driver Payroll", "Helper Payroll", "Sanjeev Driver Hisab"],
    },
    "Milk Distribution Cost": {
        "Salary Cost": ["Delivery Person Payroll", "Distribution Supervisor Payroll", "Customer Support Caller Payroll", "Anshu Hostel Feed Payroll"],
        "Rent of Hub": ["Rent Hub"],
        "Distribution Cost - Vehicle": ["Fuel - Delivery", "Petrol", "Diesel - Delivery", "Battery Swapping Charges", "Rent for Vehicle", "Bike Repair", "Delivery Bag Provision"],
        "Customer Adjustments": ["Recharge Bonus to Customer", "Adjustment for Customer", "Discount", "Offers & Discount"],
    },
    "Marketing": {
        "Advertisement": ["Advertisement", "Facebook Advertisement", "Promotional Discounts/Cashback/Referral Bonus", "LinkedIn"],
        "Marketing Materials": ["Designing Cost", "Printing Cost", "Marketing Materials - Others"],
        "Salary": ["Marketing Manager Payroll", "Marketing Executive Payroll"],
    },
    "Software Expenses/IT": {
        "SAAS / Software": ["ERP Charges", "SAAS - Milk Master/Milk Man", "MS Office", "Microsoft 365", "Google Drive", "Web Hosting/ Website", "Website Development and Maintenance", "Accounting Software", "HR Payroll Software", "Software/Software Related"],
    },
    "Office Expenses": {
        "Salary": ["Director Salary", "Top Management Salary", "Support Staff Salary", "Salary"],
        "Office Rent": ["Office Rent & Premises"],
        "Office Running": ["Stationery", "Stationery and Printing", "Travelling", "Mobile Bill", "SMS/WhatsApp Recharge", "PG Charges", "Courier Charges", "CCTV Repair", "Computer Repair"],
        "Professional Fees": ["Professional Fees", "Legal Charges/Services", "Affidavit Exp/Legal Charges"],
    },
    "Finance & Bank Charges": {
        "Banking / Statutory": ["Banking Charges", "Challan"],
    },
    "Other / General": {
        "Other": ["Donation", "Mango Purchase", "Battery", "General"],
    },
}

COST_CENTRES = [
    "Farm / Production", "Milk Procurement", "Processing Unit",
    "Milk Processing and Packaging", "Milk Transportation", "Milk Distribution",
    "Marketing", "IT / Software", "Office / Admin", "Finance / Legal",
    "Rajendra Nagar Hub", "Danapur Hub", "Patna General", "Common / Shared",
    "Other / Admin",
]
PAYMENT_MODES = ["Cash", "UPI", "Bank Transfer", "Auto-debit", "Cheque", "Card"]
PAYMENT_STATUSES = ["Paid", "Partly Paid", "Unpaid", "Accrued / Provision", "Reimbursed", "Not Applicable"]
PRODUCT_LINES = ["Cow Milk - Pouch", "Cow Milk - Bottle", "Buffalo Milk", "Paneer", "Curd", "Buttermilk", "Ghee", "Common / Shared"]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class VendorCreate(BaseModel):
    name: str
    contact: Optional[str] = None

class Vendor(VendorCreate):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: str = Field(default_factory=now_iso)

class CategoryCreate(BaseModel):
    main_category: str
    sub_category: str
    detailed_head: str

class Category(CategoryCreate):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    custom: bool = True
    created_at: str = Field(default_factory=now_iso)

class BudgetCreate(BaseModel):
    month: str  # YYYY-MM
    main_category: str
    amount: float

class Budget(BudgetCreate):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: str = Field(default_factory=now_iso)
    updated_at: str = Field(default_factory=now_iso)

class ExpenseBase(BaseModel):
    date: str  # YYYY-MM-DD
    vendor: str
    main_category: str
    sub_category: Optional[str] = None
    detailed_head: Optional[str] = None
    amount: float = 0.0
    payment_status: str = "Unpaid"
    payment_mode: Optional[str] = None
    paid_amount: float = 0.0
    # optional business detail
    voucher_no: Optional[str] = None
    cost_centre: Optional[str] = None
    product_line: Optional[str] = None
    gst_available: bool = False
    taxable_amount: Optional[float] = None
    gst_percent: Optional[float] = None
    gst_amount: Optional[float] = None
    tds: Optional[float] = None
    due_date: Optional[str] = None
    payment_date: Optional[str] = None
    bank_cash_account: Optional[str] = None
    remarks: Optional[str] = None
    recurring: bool = False
    recurring_id: Optional[str] = None

class ExpenseCreate(ExpenseBase):
    pass

class PaymentEntry(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    amount: float
    mode: Optional[str] = None
    date: str
    note: Optional[str] = None
    created_at: str = Field(default_factory=now_iso)

class Expense(ExpenseBase):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    month: str = ""
    outstanding: float = 0.0
    payments: List[PaymentEntry] = []
    created_at: str = Field(default_factory=now_iso)
    updated_at: str = Field(default_factory=now_iso)
    deleted_at: Optional[str] = None

class PaymentRecord(BaseModel):
    paid_amount: float
    payment_mode: Optional[str] = None
    payment_date: Optional[str] = None
    note: Optional[str] = None

class Recurring(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    vendor: str
    main_category: str
    sub_category: Optional[str] = None
    detailed_head: Optional[str] = None
    amount: float = 0.0
    cost_centre: Optional[str] = None
    product_line: Optional[str] = None
    tds: Optional[float] = None
    remarks: Optional[str] = None
    day_of_month: int = 1
    active: bool = True
    last_generated: Optional[str] = None
    created_at: str = Field(default_factory=now_iso)


def _derive(expense: dict) -> dict:
    """Compute month, outstanding and reconcile payment_status."""
    amount = float(expense.get("amount") or 0)
    paid = float(expense.get("paid_amount") or 0)
    outstanding = round(amount - paid, 2)
    if outstanding < 0:
        outstanding = 0.0
    expense["outstanding"] = outstanding
    date = expense.get("date") or ""
    expense["month"] = date[:7] if len(date) >= 7 else ""
    status = expense.get("payment_status")
    if status not in ("Not Applicable", "Accrued / Provision", "Reimbursed"):
        if paid <= 0:
            expense["payment_status"] = "Unpaid"
        elif paid >= amount and amount > 0:
            expense["payment_status"] = "Paid"
        else:
            expense["payment_status"] = "Partly Paid"
    return expense


# ---------------------------------------------------------------------------
# Masters
# ---------------------------------------------------------------------------
@api_router.get("/")
async def root():
    return {"message": "Paavni Bookkeeping API"}

@api_router.get("/masters")
async def get_masters():
    categories = []
    for main, subs in CATEGORY_HIERARCHY.items():
        sub_list = []
        for sub, heads in subs.items():
            sub_list.append({"sub_category": sub, "detailed_heads": heads})
        categories.append({"main_category": main, "sub_categories": sub_list})
    custom = await db.categories.find({"deleted_at": None}).to_list(1000)
    return {
        "categories": categories,
        "custom_categories": [Category(**c).dict() for c in custom],
        "cost_centres": COST_CENTRES,
        "payment_modes": PAYMENT_MODES,
        "payment_statuses": PAYMENT_STATUSES,
        "product_lines": PRODUCT_LINES,
    }


# ---------------------------------------------------------------------------
# Vendors
# ---------------------------------------------------------------------------
@api_router.get("/vendors", response_model=List[Vendor])
async def list_vendors():
    docs = await db.vendors.find({"deleted_at": None}).sort("name", 1).to_list(1000)
    return [Vendor(**d) for d in docs]

@api_router.post("/vendors", response_model=Vendor)
async def create_vendor(payload: VendorCreate):
    vendor = Vendor(**payload.dict())
    doc = vendor.dict()
    doc["deleted_at"] = None
    await db.vendors.insert_one(doc)
    return vendor

@api_router.put("/vendors/{vendor_id}", response_model=Vendor)
async def update_vendor(vendor_id: str, payload: VendorCreate):
    existing = await db.vendors.find_one({"id": vendor_id, "deleted_at": None})
    if not existing:
        raise HTTPException(status_code=404, detail="Vendor not found")
    await db.vendors.update_one({"id": vendor_id}, {"$set": payload.dict()})
    updated = await db.vendors.find_one({"id": vendor_id})
    return Vendor(**updated)

@api_router.delete("/vendors/{vendor_id}")
async def delete_vendor(vendor_id: str):
    await db.vendors.update_one({"id": vendor_id}, {"$set": {"deleted_at": now_iso()}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Categories (custom additions)
# ---------------------------------------------------------------------------
@api_router.get("/categories", response_model=List[Category])
async def list_categories():
    docs = await db.categories.find({"deleted_at": None}).to_list(1000)
    return [Category(**d) for d in docs]

@api_router.post("/categories", response_model=Category)
async def create_category(payload: CategoryCreate):
    cat = Category(**payload.dict())
    doc = cat.dict()
    doc["deleted_at"] = None
    await db.categories.insert_one(doc)
    return cat

@api_router.delete("/categories/{category_id}")
async def delete_category(category_id: str):
    await db.categories.update_one({"id": category_id}, {"$set": {"deleted_at": now_iso()}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Expenses
# ---------------------------------------------------------------------------
@api_router.post("/expenses", response_model=Expense)
async def create_expense(payload: ExpenseCreate):
    expense = Expense(**payload.dict())
    doc = expense.dict()
    doc = _derive(doc)
    # Seed the payment ledger if an initial amount was already paid.
    if float(doc.get("paid_amount") or 0) > 0:
        doc["payments"] = [PaymentEntry(
            amount=float(doc["paid_amount"]),
            mode=doc.get("payment_mode"),
            date=doc.get("payment_date") or doc.get("date"),
            note="Initial payment",
        ).dict()]
    # If flagged recurring, create a monthly template linked to this expense.
    if payload.recurring and not payload.recurring_id:
        try:
            day = int((doc.get("date") or "")[8:10]) or 1
        except ValueError:
            day = 1
        tmpl = Recurring(
            vendor=doc["vendor"],
            main_category=doc["main_category"],
            sub_category=doc.get("sub_category"),
            detailed_head=doc.get("detailed_head"),
            amount=float(doc.get("amount") or 0),
            cost_centre=doc.get("cost_centre"),
            product_line=doc.get("product_line"),
            tds=doc.get("tds"),
            remarks=doc.get("remarks"),
            day_of_month=day,
            last_generated=doc.get("month"),
        ).dict()
        tmpl["deleted_at"] = None
        await db.recurring.insert_one(tmpl)
        doc["recurring_id"] = tmpl["id"]
        doc["recurring"] = True
    doc["deleted_at"] = None
    await db.expenses.insert_one(doc)
    saved = await db.expenses.find_one({"id": expense.id})
    saved.pop("_id", None)
    return Expense(**saved)


async def ensure_recurring():
    """Auto-create the current month's instance for each active recurring template."""
    month = now_iso()[:7]
    templates = await db.recurring.find({"deleted_at": None, "active": True}).to_list(1000)
    for t in templates:
        if t.get("last_generated") == month:
            continue
        exists = await db.expenses.find_one({"recurring_id": t["id"], "month": month, "deleted_at": None})
        if not exists:
            y, m = int(month[:4]), int(month[5:7])
            last_day = calendar.monthrange(y, m)[1]
            day = min(int(t.get("day_of_month") or 1), last_day)
            date = f"{month}-{day:02d}"
            exp = Expense(
                date=date,
                vendor=t["vendor"],
                main_category=t["main_category"],
                sub_category=t.get("sub_category"),
                detailed_head=t.get("detailed_head"),
                amount=float(t.get("amount") or 0),
                payment_status="Unpaid",
                cost_centre=t.get("cost_centre"),
                product_line=t.get("product_line"),
                tds=t.get("tds"),
                remarks=t.get("remarks"),
                recurring=True,
                recurring_id=t["id"],
            )
            edoc = _derive(exp.dict())
            edoc["deleted_at"] = None
            await db.expenses.insert_one(edoc)
        await db.recurring.update_one({"id": t["id"]}, {"$set": {"last_generated": month}})

@api_router.get("/expenses", response_model=List[Expense])
async def list_expenses(
    month: Optional[str] = None,
    payment_status: Optional[str] = None,
    main_category: Optional[str] = None,
    vendor: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    q: Optional[str] = None,
):
    await ensure_recurring()
    query = {"deleted_at": None}
    if month:
        query["month"] = month
    if payment_status:
        query["payment_status"] = payment_status
    if main_category:
        query["main_category"] = main_category
    if vendor:
        query["vendor"] = vendor
    if date_from or date_to:
        date_q = {}
        if date_from:
            date_q["$gte"] = date_from
        if date_to:
            date_q["$lte"] = date_to
        query["date"] = date_q
    if q:
        regex = {"$regex": q, "$options": "i"}
        query["$or"] = [
            {"vendor": regex},
            {"main_category": regex},
            {"sub_category": regex},
            {"detailed_head": regex},
            {"product_line": regex},
            {"remarks": regex},
        ]
    docs = await db.expenses.find(query).sort("date", -1).to_list(2000)
    out = []
    for d in docs:
        d.pop("_id", None)
        out.append(Expense(**d))
    return out

@api_router.get("/expenses/{expense_id}", response_model=Expense)
async def get_expense(expense_id: str):
    doc = await db.expenses.find_one({"id": expense_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Expense not found")
    doc.pop("_id", None)
    return Expense(**doc)

@api_router.put("/expenses/{expense_id}", response_model=Expense)
async def update_expense(expense_id: str, payload: ExpenseCreate):
    existing = await db.expenses.find_one({"id": expense_id, "deleted_at": None})
    if not existing:
        raise HTTPException(status_code=404, detail="Expense not found")
    update = payload.dict()
    update["updated_at"] = now_iso()
    merged = {**existing, **update}
    merged = _derive(merged)
    await db.expenses.update_one({"id": expense_id}, {"$set": {
        **update,
        "month": merged["month"],
        "outstanding": merged["outstanding"],
        "payment_status": merged["payment_status"],
    }})
    updated = await db.expenses.find_one({"id": expense_id})
    updated.pop("_id", None)
    return Expense(**updated)

@api_router.post("/expenses/{expense_id}/payment", response_model=Expense)
async def record_payment(expense_id: str, payload: PaymentRecord):
    existing = await db.expenses.find_one({"id": expense_id, "deleted_at": None})
    if not existing:
        raise HTTPException(status_code=404, detail="Expense not found")
    if float(payload.paid_amount) <= 0:
        raise HTTPException(status_code=400, detail="Payment amount must be greater than zero")
    pay_date = payload.payment_date or now_iso()[:10]
    entry = PaymentEntry(
        amount=float(payload.paid_amount),
        mode=payload.payment_mode,
        date=pay_date,
        note=payload.note,
    ).dict()
    new_paid = round(float(existing.get("paid_amount") or 0) + float(payload.paid_amount), 2)
    update = {
        "paid_amount": new_paid,
        "updated_at": now_iso(),
        "payment_date": pay_date,
    }
    if payload.payment_mode:
        update["payment_mode"] = payload.payment_mode
    merged = {**existing, **update}
    merged = _derive(merged)
    update["month"] = merged["month"]
    update["outstanding"] = merged["outstanding"]
    update["payment_status"] = merged["payment_status"]
    await db.expenses.update_one(
        {"id": expense_id},
        {"$set": update, "$push": {"payments": entry}},
    )
    updated = await db.expenses.find_one({"id": expense_id})
    updated.pop("_id", None)
    return Expense(**updated)

@api_router.delete("/expenses/{expense_id}")
async def delete_expense(expense_id: str):
    await db.expenses.update_one({"id": expense_id}, {"$set": {"deleted_at": now_iso()}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@api_router.get("/dashboard")
async def dashboard(month: Optional[str] = None):
    await ensure_recurring()
    query = {"deleted_at": None}
    if month:
        query["month"] = month
    docs = await db.expenses.find(query).to_list(5000)
    total = 0.0
    paid = 0.0
    outstanding = 0.0
    by_category = {}
    by_status = {}
    for d in docs:
        amt = float(d.get("amount") or 0)
        p = float(d.get("paid_amount") or 0)
        o = float(d.get("outstanding") or 0)
        total += amt
        paid += p
        outstanding += o
        cat = d.get("main_category") or "Other / General"
        by_category[cat] = by_category.get(cat, 0.0) + amt
        st = d.get("payment_status") or "Unpaid"
        entry = by_status.setdefault(st, {"count": 0, "amount": 0.0})
        entry["count"] += 1
        entry["amount"] += amt

    cat_list = sorted(
        [{"category": k, "amount": round(v, 2)} for k, v in by_category.items()],
        key=lambda x: x["amount"], reverse=True,
    )
    status_list = [{"status": k, "count": v["count"], "amount": round(v["amount"], 2)} for k, v in by_status.items()]

    months = await db.expenses.distinct("month", {"deleted_at": None})
    months = sorted([m for m in months if m], reverse=True)

    # Budget vs actual (only meaningful for a specific month)
    budget_block = {"has_budgets": False, "total_budget": 0.0, "total_spent": 0.0, "by_category": []}
    if month:
        bdocs = await db.budgets.find({"month": month, "deleted_at": None}).to_list(1000)
        if bdocs:
            b_list = []
            total_budget = 0.0
            total_spent_budgeted = 0.0
            for b in bdocs:
                cat = b["main_category"]
                amt = float(b.get("amount") or 0)
                spent = round(float(by_category.get(cat, 0.0)), 2)
                total_budget += amt
                total_spent_budgeted += spent
                b_list.append({
                    "category": cat,
                    "budget": round(amt, 2),
                    "spent": spent,
                    "remaining": round(amt - spent, 2),
                    "pct": round(spent / amt * 100, 1) if amt > 0 else 0.0,
                })
            b_list.sort(key=lambda x: x["pct"], reverse=True)
            budget_block = {
                "has_budgets": True,
                "total_budget": round(total_budget, 2),
                "total_spent": round(total_spent_budgeted, 2),
                "by_category": b_list,
            }

    return {
        "month": month,
        "total": round(total, 2),
        "paid": round(paid, 2),
        "outstanding": round(outstanding, 2),
        "count": len(docs),
        "by_category": cat_list,
        "by_status": status_list,
        "months": months,
        "budget": budget_block,
    }


# ---------------------------------------------------------------------------
# Budgets
# ---------------------------------------------------------------------------
@api_router.get("/budgets", response_model=List[Budget])
async def list_budgets(month: Optional[str] = None):
    query = {"deleted_at": None}
    if month:
        query["month"] = month
    docs = await db.budgets.find(query).to_list(1000)
    out = []
    for d in docs:
        d.pop("_id", None)
        out.append(Budget(**d))
    return out

@api_router.post("/budgets", response_model=Budget)
async def upsert_budget(payload: BudgetCreate):
    existing = await db.budgets.find_one({
        "month": payload.month,
        "main_category": payload.main_category,
        "deleted_at": None,
    })
    if existing:
        await db.budgets.update_one(
            {"id": existing["id"]},
            {"$set": {"amount": float(payload.amount), "updated_at": now_iso()}},
        )
        updated = await db.budgets.find_one({"id": existing["id"]})
        updated.pop("_id", None)
        return Budget(**updated)
    budget = Budget(**payload.dict())
    doc = budget.dict()
    doc["deleted_at"] = None
    await db.budgets.insert_one(doc)
    return budget

@api_router.delete("/budgets/{budget_id}")
async def delete_budget(budget_id: str):
    await db.budgets.update_one({"id": budget_id}, {"$set": {"deleted_at": now_iso()}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Vendor dues (vendor-wise outstanding)
# ---------------------------------------------------------------------------
@api_router.get("/vendors/dues")
async def vendor_dues():
    docs = await db.expenses.find({"deleted_at": None}).to_list(5000)
    agg = {}
    for d in docs:
        v = d.get("vendor") or "Unknown"
        e = agg.setdefault(v, {"vendor": v, "total": 0.0, "paid": 0.0, "outstanding": 0.0, "bills": 0, "open_bills": 0})
        amt = float(d.get("amount") or 0)
        paid = float(d.get("paid_amount") or 0)
        out = float(d.get("outstanding") or 0)
        e["total"] += amt
        e["paid"] += paid
        e["outstanding"] += out
        e["bills"] += 1
        if out > 0:
            e["open_bills"] += 1
    result = [
        {**v, "total": round(v["total"], 2), "paid": round(v["paid"], 2), "outstanding": round(v["outstanding"], 2)}
        for v in agg.values()
    ]
    result.sort(key=lambda x: x["outstanding"], reverse=True)
    total_outstanding = round(sum(v["outstanding"] for v in result), 2)
    return {"vendors": result, "total_outstanding": total_outstanding}


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
@api_router.get("/export")
async def export_expenses(month: Optional[str] = None, vendor: Optional[str] = None):
    query = {"deleted_at": None}
    if month:
        query["month"] = month
    if vendor:
        query["vendor"] = vendor
    docs = await db.expenses.find(query).sort("date", 1).to_list(5000)

    wb = Workbook()
    ws = wb.active
    ws.title = (vendor or month or "All")[:31]

    headers = [
        "Date", "Vendor", "Main Category", "Sub Category", "Detailed Head",
        "Cost Centre", "Product Line", "Taxable", "GST %", "GST Amount",
        "TDS", "Amount", "Paid", "Outstanding", "Status", "Payment Mode", "Remarks",
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="2A5C43", end_color="2A5C43", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_amt = total_paid = total_out = 0.0
    for d in docs:
        total_amt += float(d.get("amount") or 0)
        total_paid += float(d.get("paid_amount") or 0)
        total_out += float(d.get("outstanding") or 0)
        ws.append([
            d.get("date", ""),
            d.get("vendor", ""),
            d.get("main_category", ""),
            d.get("sub_category", "") or "",
            d.get("detailed_head", "") or "",
            d.get("cost_centre", "") or "",
            d.get("product_line", "") or "",
            d.get("taxable_amount") or "",
            d.get("gst_percent") or "",
            d.get("gst_amount") or "",
            d.get("tds") or "",
            float(d.get("amount") or 0),
            float(d.get("paid_amount") or 0),
            float(d.get("outstanding") or 0),
            d.get("payment_status", ""),
            d.get("payment_mode", "") or "",
            d.get("remarks", "") or "",
        ])

    total_row = ["", "", "", "", "", "", "", "", "", "", "TOTAL",
                 round(total_amt, 2), round(total_paid, 2), round(total_out, 2), "", "", ""]
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    widths = [12, 22, 20, 18, 22, 16, 16, 12, 8, 12, 10, 12, 12, 12, 12, 14, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    slug = vendor.replace(" ", "-").lower() if vendor else (month or "all")
    fname = f"paavni-expenses-{slug}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ---------------------------------------------------------------------------
# Recurring expense templates
# ---------------------------------------------------------------------------
@api_router.get("/recurring", response_model=List[Recurring])
async def list_recurring():
    docs = await db.recurring.find({"deleted_at": None}).sort("vendor", 1).to_list(1000)
    out = []
    for d in docs:
        d.pop("_id", None)
        out.append(Recurring(**d))
    return out

@api_router.delete("/recurring/{recurring_id}")
async def delete_recurring(recurring_id: str):
    await db.recurring.update_one({"id": recurring_id}, {"$set": {"deleted_at": now_iso(), "active": False}})
    return {"ok": True}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
